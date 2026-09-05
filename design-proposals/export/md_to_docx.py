# -*- coding: utf-8 -*-
"""将 v1.1 架构文档 markdown 导出为 Word docx（python-docx）。
用法: python md_to_docx.py [输入.md] [输出.docx]
"""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SRC = sys.argv[1] if len(sys.argv) > 1 else r"D:\Tools\deepseek-harness\design-proposals\industrial-automation-ai-agent-closed-loop-architecture.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else r"D:\Tools\deepseek-harness\design-proposals\工业自动化AI-Agent闭环架构-v1.1.docx"

TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
DIAG_NAMES = ["三闭环总览", "分层架构", "闭环时序示例（需求到现场交付）"]

def set_east_asia(style_or_rpr, font="微软雅黑"):
    el = style_or_rpr
    rfonts = el.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        el.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsia"), "微软雅黑")
    rPr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rPr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)

def add_runs(paragraph, text, size=None, bold_all=False):
    for seg in TOKEN.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = paragraph.add_run(seg[2:-2])
            r.bold = True
        elif seg.startswith("`") and seg.endswith("`"):
            r = paragraph.add_run(seg[1:-1])
            r.font.name = "Consolas"
            set_east_asia(r._element.get_or_add_rPr(), "微软雅黑")
        elif seg.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", seg)
            if m:
                add_hyperlink(paragraph, m.group(1), m.group(2))
                continue
            r = paragraph.add_run(seg)
        else:
            r = paragraph.add_run(seg)
        if bold_all:
            r.bold = True
        if size:
            r.font.size = Pt(size)
    return paragraph

def shade_paragraph(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def shade_cell(cell, fill="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)

for sname in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
    st = doc.styles[sname]
    st.font.name = "Calibri"
    set_east_asia(st.element.get_or_add_rPr(), "微软雅黑")
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
doc.styles["Title"].font.size = Pt(20)
doc.styles["Heading 1"].font.size = Pt(15)
doc.styles["Heading 2"].font.size = Pt(12.5)
doc.styles["Heading 3"].font.size = Pt(11.5)
doc.styles["Heading 1"].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
doc.styles["Heading 2"].font.color.rgb = RGBColor(0x2E, 0x5A, 0x9E)
doc.styles["Heading 3"].font.color.rgb = RGBColor(0x2E, 0x5A, 0x9E)

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i = 0
n = len(lines)
diag_idx = 0
while i < n:
    line = lines[i]
    stripped = line.strip()
    if not stripped:
        i += 1
        continue
    if stripped.startswith("```mermaid"):
        name = DIAG_NAMES[diag_idx] if diag_idx < len(DIAG_NAMES) else "架构图"
        diag_idx += 1
        cap = doc.add_paragraph()
        r = cap.add_run("【图：%s】（Mermaid 源码，可在支持 Mermaid 的编辑器中渲染）" % name)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        i += 1
        while i < n and not lines[i].strip().startswith("```"):
            p = doc.add_paragraph()
            r = p.add_run(lines[i])
            r.font.name = "Consolas"
            r.font.size = Pt(8)
            set_east_asia(r._element.get_or_add_rPr(), "微软雅黑")
            p.paragraph_format.space_after = Pt(0)
            shade_paragraph(p)
            i += 1
        i += 1
        continue
    if stripped.startswith("|"):
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            rows.append(lines[i])
            i += 1
        grid = []
        for row in rows:
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                continue
            grid.append(cells)
        if grid:
            ncols = max(len(r) for r in grid)
            table = doc.add_table(rows=len(grid), cols=ncols)
            table.style = "Table Grid"
            for ri, row in enumerate(grid):
                for ci in range(ncols):
                    cell = table.cell(ri, ci)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    text = row[ci] if ci < len(row) else ""
                    if ri == 0:
                        shade_cell(cell)
                    add_runs(cell.paragraphs[0], text, size=9.5, bold_all=(ri == 0))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    if stripped.startswith("> "):
        p = doc.add_paragraph()
        add_runs(p, stripped[2:], size=10)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            r.italic = True
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(4)
        i += 1
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
    if m:
        level = len(m.group(1)) - 1
        doc.add_heading(m.group(2), level=level)
        i += 1
        continue
    m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
    if m:
        p = doc.add_paragraph()
        add_runs(p, "%s. %s" % (m.group(1), m.group(2)))
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        i += 1
        continue
    m = re.match(r"^-\s+(.*)$", stripped)
    if m:
        p = doc.add_paragraph()
        add_runs(p, "• " + m.group(1))
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        i += 1
        continue
    p = doc.add_paragraph()
    add_runs(p, stripped)
    i += 1

doc.save(OUT)
print("DOCX SAVED:", OUT)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))

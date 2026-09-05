# -*- coding: utf-8 -*-
"""读取 .pptx：逐页提取文本与表格，输出到 txt/md 或打印摘要。
用法: python pptx_extract.py <pptx路径> [输出.md]
"""
import sys
from pptx import Presentation

def dump(path, out_md=None):
    prs = Presentation(path)
    w, h = prs.slide_width, prs.slide_height
    lines = []
    lines.append("# PPTX 提取: %s" % path)
    lines.append("> 页数: %d ｜ 尺寸: %.2f x %.2f 英寸 (%d x %d pt)" %
                 (len(prs.slides), w / 914400, h / 914400, w / 12700, h / 12700))
    for idx, slide in enumerate(prs.slides, 1):
        lines.append("")
        lines.append("## 第 %d 页" % idx)
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
                lines.append("- 文本框: " + txt.replace("\n", " ⏎ "))
            if sh.has_table:
                tbl = sh.table
                lines.append("- 表格 (%d 行 x %d 列):" % (len(tbl.rows), len(tbl.columns)))
                for r in tbl.rows:
                    cells = [c.text.replace("\n", " ") for c in r.cells]
                    lines.append("  | " + " | ".join(cells) + " |")
    text = "\n".join(lines)
    if out_md:
        with open(out_md, "w", encoding="utf-8") as f:
            f.write(text)
        print("SAVED:", out_md, "chars:", len(text))
    else:
        print(text)
    return len(prs.slides)

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else None
    if not src:
        print("usage: python pptx_extract.py <pptx路径> [输出.md]")
        sys.exit(2)
    out = sys.argv[2] if len(sys.argv) > 2 else None
    dump(src, out)
